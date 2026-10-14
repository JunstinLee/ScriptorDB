from __future__ import annotations

import asyncio
import csv
import glob as glob_mod
import json
import os
from typing import Any

from pydantic_ai import RunContext

from config.settings import Settings
from tools.errors import _to_tool_error
from tools.pdf.tools import _read_pdf
from tools.tool_decorators import db_tool
from tools.tool_result import ToolErrorInfo, ToolResult
from tools.validators import validate_file_path


def _read_delimited(
    ctx: RunContext[Settings],
    filepath: str,
    preview_rows: int = 10,
    encoding: str = "utf-8",
    return_full: bool = False,
    max_rows: int | None = None,
) -> ToolResult:
    if not os.path.isfile(filepath):
        return ToolResult(
            success=False,
            error=ToolErrorInfo(
                category="resource_not_found",
                message=f"File not found: {filepath}",
            ),
        )

    try:
        with open(filepath, "r", encoding=encoding, newline="") as f:
            reader = csv.reader(f)
            try:
                headers = next(reader)
            except StopIteration:
                headers = []
            preview: list[list[str]] = []
            rows_data: list[list[str]] = []
            row_count = 0
            truncated = False
            for row in reader:
                row_count += 1
                if len(preview) < preview_rows:
                    preview.append(row)
                if return_full:
                    if max_rows is None or len(rows_data) < max_rows:
                        rows_data.append(row)
                    else:
                        truncated = True

        data: dict[str, Any] = {
            "file": filepath,
            "columns": headers,
            "rows": row_count,
            "preview": preview,
        }
        output = (
            f"Read {os.path.basename(filepath)}: {row_count} row"
            f"{'s' if row_count != 1 else ''}, {len(headers)} column"
            f"{'s' if len(headers) != 1 else ''}"
        )
        if return_full:
            data["data"] = rows_data
            data["truncated"] = truncated
            if truncated:
                output += f" (returned {len(rows_data)} of {row_count} rows)"

        return ToolResult(success=True, output=output, data=data)
    except FileNotFoundError:
        return ToolResult(
            success=False,
            error=ToolErrorInfo(
                category="resource_not_found",
                message=f"File not found: {filepath}",
            ),
        )
    except Exception as e:
        return _to_tool_error(e)


@db_tool(name="write_csv", category="write", timeout=30, requires_approval=True, validator=validate_file_path)
def write_csv(
    ctx: RunContext[Settings],
    filepath: str,
    columns: list[str],
    data: list[list[str]],
    encoding: str = "utf-8",
) -> ToolResult:
    try:
        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)
        with open(filepath, "w", encoding=encoding, newline="") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(data)

        return ToolResult(
            success=True,
            output=f"Wrote {len(data)} row{'s' if len(data) != 1 else ''} to {os.path.basename(filepath)}",
            data={"file": filepath, "rows_written": len(data)},
        )
    except Exception as e:
        return _to_tool_error(e)


def _read_excel(
    ctx: RunContext[Settings],
    filepath: str,
    sheet_name: str | int = 0,
    header_row: int = 1,
    preview_rows: int = 10,
    return_full: bool = False,
    max_rows: int | None = None,
) -> ToolResult:
    if not os.path.isfile(filepath):
        return ToolResult(
            success=False,
            error=ToolErrorInfo(
                category="resource_not_found",
                message=f"File not found: {filepath}",
            ),
        )

    try:
        from openpyxl import load_workbook
    except ImportError:
        return ToolResult(
            success=False,
            error=ToolErrorInfo(
                category="parameter_error",
                message="openpyxl is not installed. Run: uv sync",
            ),
        )

    try:
        wb = load_workbook(filepath, data_only=True, read_only=True)
        if isinstance(sheet_name, int):
            if sheet_name < 0 or sheet_name >= len(wb.worksheets):
                return ToolResult(
                    success=False,
                    error=ToolErrorInfo(
                        category="parameter_error",
                        message=f"Sheet index {sheet_name} out of range",
                    ),
                )
            ws = wb.worksheets[sheet_name]
        else:
            if sheet_name not in wb.sheetnames:
                return ToolResult(
                    success=False,
                    error=ToolErrorInfo(
                        category="parameter_error",
                        message=f"Sheet '{sheet_name}' not found",
                    ),
                )
            ws = wb[sheet_name]

        rows_iter = ws.iter_rows(min_row=header_row, values_only=True)
        try:
            headers = list(next(rows_iter))
        except StopIteration:
            headers = []

        preview: list[list[Any | None]] = []
        rows_data: list[list[Any | None]] = []
        row_count = 0
        truncated = False
        for row in rows_iter:
            row_count += 1
            row_values = [str(v) if v is not None else None for v in row]
            if len(preview) < preview_rows:
                preview.append(row_values)
            if return_full:
                if max_rows is None or len(rows_data) < max_rows:
                    rows_data.append(row_values)
                else:
                    truncated = True

        wb.close()

        data: dict[str, Any] = {
            "file": filepath,
            "sheet": ws.title,
            "columns": headers,
            "rows": row_count,
            "preview": preview,
        }
        output = (
            f"Read {os.path.basename(filepath)}: sheet {ws.title}, {row_count} row"
            f"{'s' if row_count != 1 else ''}, {len(headers)} column"
            f"{'s' if len(headers) != 1 else ''}"
        )
        if return_full:
            data["data"] = rows_data
            data["truncated"] = truncated
            if truncated:
                output += f" (returned {len(rows_data)} of {row_count} rows)"

        return ToolResult(success=True, output=output, data=data)
    except Exception as e:
        return _to_tool_error(e)


@db_tool(name="read_file", timeout=60, validator=validate_file_path)
async def read_file(
    ctx: RunContext[Settings],
    filepath: str,
    sheet_name: str | int = 0,
    header_row: int = 1,
    preview_rows: int = 10,
    return_full: bool = False,
    max_rows: int | None = None,
    max_size_kb: int = 50,
    max_chars: int = 50000,
    encoding: str = "utf-8",
) -> ToolResult:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".csv":
        return await asyncio.to_thread(
            _read_delimited,
            ctx,
            filepath,
            preview_rows=preview_rows,
            encoding=encoding,
            return_full=return_full,
            max_rows=max_rows,
        )
    if ext in {".xlsx", ".xls"}:
        return await asyncio.to_thread(
            _read_excel,
            ctx,
            filepath,
            sheet_name=sheet_name,
            header_row=header_row,
            preview_rows=preview_rows,
            return_full=return_full,
            max_rows=max_rows,
        )
    if ext == ".pdf":
        return await _read_pdf(ctx, filepath, max_chars=max_chars)
    return _read_text(ctx, filepath, max_size_kb=max_size_kb, encoding=encoding)


def _read_text(
    ctx: RunContext[Settings],
    filepath: str,
    max_size_kb: int = 50,
    encoding: str = "utf-8",
) -> ToolResult:
    if not os.path.isfile(filepath):
        return ToolResult(
            success=False,
            error=ToolErrorInfo(
                category="resource_not_found",
                message=f"File not found: {filepath}",
            ),
        )

    try:
        size = os.path.getsize(filepath)
        truncated = size > max_size_kb * 1024

        with open(filepath, "r", encoding=encoding) as f:
            content = f.read(max_size_kb * 1024) if truncated else f.read()

        ext = os.path.splitext(filepath)[1].lower()
        extra: dict = {"file": filepath, "size_bytes": size, "truncated": truncated}

        if ext == ".json":
            try:
                parsed = json.loads(content)
                extra["parsed"] = True
                extra["json_keys"] = list(parsed.keys()) if isinstance(parsed, dict) else None
                output = f"JSON file {os.path.basename(filepath)}{' (truncated)' if truncated else ''}"
                return ToolResult(success=True, output=output, data=extra)
            except json.JSONDecodeError:
                pass

        output = f"File {os.path.basename(filepath)}: {size} bytes{' (truncated)' if truncated else ''}"
        return ToolResult(success=True, output=output, data=extra | {"content": content})
    except Exception as e:
        return _to_tool_error(e)


@db_tool(name="write_file", category="write", timeout=10, requires_approval=True, validator=validate_file_path)
def write_file(
    ctx: RunContext[Settings],
    filepath: str,
    content: str,
    encoding: str = "utf-8",
) -> ToolResult:
    try:
        os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)
        with open(filepath, "w", encoding=encoding) as f:
            f.write(content)

        return ToolResult(
            success=True,
            output=f"Wrote {len(content)} character{'s' if len(content) != 1 else ''} to {os.path.basename(filepath)}",
            data={"file": filepath, "chars_written": len(content)},
        )
    except Exception as e:
        return _to_tool_error(e)


@db_tool(name="list_files", timeout=5)
def list_files(
    ctx: RunContext[Settings],
    directory: str = ".",
    pattern: str = "*",
) -> ToolResult:
    try:
        search_path = os.path.join(directory, pattern)
        files = sorted(glob_mod.glob(search_path))
        result_files: list[dict] = []
        for f in files[:100]:
            try:
                st = os.stat(f)
                result_files.append({
                    "name": f,
                    "size_bytes": st.st_size,
                    "is_dir": os.path.isdir(f),
                })
            except OSError:
                result_files.append({"name": f, "size_bytes": 0, "is_dir": False})

        return ToolResult(
            success=True,
            output=f"{len(files)} file{'s' if len(files) != 1 else ''}/director{'ies' if len(files) != 1 else 'y'} matching {pattern}",
            data={
                "directory": directory,
                "pattern": pattern,
                "count": min(len(files), 100),
                "files": result_files,
            },
        )
    except Exception as e:
        return _to_tool_error(e)
